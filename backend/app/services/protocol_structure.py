"""Word 方案正文结构；只解析内容和原文位置，不承载字段业务规则。"""
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .protocol_cell_images import cell_images


def _blocks(parent):
    for child in parent:
        if child.tag in {qn('w:p'), qn('w:tbl')}:
            yield child
        elif child.tag in {qn('w:sdt'), qn('w:sdtContent'), qn('w:customXml')}:
            yield from _blocks(child)


def _heading_level(paragraph: Paragraph) -> int | None:
    properties = paragraph._p.pPr
    if properties is not None and properties.find(qn('w:outlineLvl')) is not None:
        value = int(properties.find(qn('w:outlineLvl')).get(qn('w:val')))
        return value if value < 9 else None
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        properties = style.element.find(qn('w:pPr'))
        outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
        if outline is not None:
            value = int(outline.get(qn('w:val')))
            return value if value < 9 else None
        style = style.base_style
    return None


def read_protocol_structure(path: Path, include_images: bool = False) -> list[dict[str, Any]]:
    document = Document(path)
    blocks = []
    headings: list[tuple[int, str]] = []
    paragraph_number = table_number = 0
    for element in _blocks(document.element.body):
        if element.tag == qn('w:p'):
            paragraph_number += 1
            paragraph = Paragraph(element, document)
            # XML 文本节点包括内容控件里的文字；标签跨 run 时仍保持连续。
            text = _paragraph_text(element)
            level = _heading_level(paragraph)
            if level is not None:
                headings = [(depth, title) for depth, title in headings if depth < level]
                headings.append((level, text))
            blocks.append({'kind': 'paragraph', 'text': text, 'level': level,
                           'paragraph': paragraph_number, 'section': ' / '.join(title for _, title in headings)})
        else:
            table_number += 1
            table = Table(element, document)
            try:
                rows = _table_rows(table, table_number)
            except ValueError as error:
                # 结构错误保留到规则选中章节时抛出，避免无关章节阻断提取。
                blocks.append({'kind': 'table', 'table': table_number, 'error': str(error),
                               'section': ' / '.join(title for _, title in headings)})
                continue
            block = {'kind': 'table', 'rows': rows, 'table': table_number,
                     'text': '\n'.join('\t'.join(row) for row in rows),
                     'section': ' / '.join(title for _, title in headings)}
            if include_images:
                try:
                    block['images'] = [[cell_images(cell, document.part) for cell in row.cells] for row in table.rows]
                except ValueError as error:
                    block['image_error'] = str(error)
            blocks.append(block)
    return blocks


def read_protocol_header_structure(path: Path) -> list[dict[str, Any]]:
    """读取方案已有页眉；字段定位条件由提取规则负责。"""
    document = Document(path)
    blocks = []
    seen_parts = set()
    table_number = 0
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            part_name = str(header.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for table in header.tables:
                table_number += 1
                try:
                    rows = _table_rows(table, table_number)
                except ValueError as error:
                    blocks.append({'kind': 'table', 'table': table_number, 'headerPart': part_name,
                                   'error': str(error), 'section': '页眉'})
                    continue
                blocks.append({'kind': 'table', 'rows': rows, 'table': table_number,
                               'headerPart': part_name, 'section': '页眉'})
    return blocks


def _table_rows(table: Table, table_number: int) -> list[list[str]]:
    rows = []
    width = len(table.columns)
    for row in table.rows:
        if row.grid_cols_before or row.grid_cols_after or len(row.cells) != width:
            raise ValueError(f'方案表格 {table_number} 存在缺省网格单元格，请补齐表格结构')
        rows.append([_cell_text(cell._tc) for cell in row.cells])
    return rows


def _paragraph_text(element) -> str:
    parts = []
    for child in element.iter():
        if child.tag == qn('w:t'):
            parts.append(child.text or '')
        elif child.tag == qn('w:tab'):
            parts.append('\t')
        elif child.tag in {qn('w:br'), qn('w:cr')}:
            parts.append('\n')
    return ''.join(parts)


def _cell_text(element) -> str:
    blocks = list(_blocks(element))
    if any(block.tag == qn('w:tbl') for block in blocks):
        raise ValueError('方案包含嵌套表格，请将待提取明细整理为独立表格')
    return '\n'.join(_paragraph_text(paragraph) for paragraph in blocks)
