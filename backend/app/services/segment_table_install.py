"""给当前草稿安装行片段表规则；不修改 Word 模板中的行列或内容。"""

import copy
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from lxml import etree

from .docx_language import write_docx_parts_atomic


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = f"{{{NS['w']}}}"


def target_table(path: Path, header: str) -> tuple[int, set[str]]:
    document = Document(path)
    matches = [(index, table) for index, table in enumerate(document.tables, 1)
               if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip() == header]
    if len(matches) != 1:
        raise ValueError(f"Word 模板中必须且只能有一张首列表头为 {header} 的表")
    index, table = matches[0]
    return index, set(table._tbl.xpath(".//w:sdtPr/w:tag/@w:val"))


def _unique_tags(path: Path, table_index: int, controls: list[dict[str, Any]]) -> None:
    if not controls:
        return
    with zipfile.ZipFile(path) as archive:
        parts = {entry.filename: (copy.copy(entry), archive.read(entry.filename))
                 for entry in archive.infolist()}
    name = "word/document.xml"
    root = etree.fromstring(parts[name][1])
    tables = root.xpath("./w:body/w:tbl", namespaces=NS)
    table = tables[table_index - 1]
    existing = set(root.xpath(".//w:sdtPr/w:tag/@w:val", namespaces=NS))
    changed = False
    for item in controls:
        rows = table.xpath("./w:tr", namespaces=NS)
        row = int(item["row"]) - 1
        column = int(item["column"]) - 1
        if row < 0 or row >= len(rows):
            raise ValueError("唯一控件的原型行超出 Word 表格范围")
        cells = rows[row].xpath("./w:tc", namespaces=NS)
        if column < 0 or column >= len(cells):
            raise ValueError("唯一控件的原型列超出 Word 表格范围")
        tag = str(item["tag"])
        nodes = cells[column].xpath(".//w:sdtPr/w:tag", namespaces=NS)
        if len(nodes) != 1:
            raise ValueError("唯一控件所在单元格必须且只能包含一个内容控件")
        current = nodes[0].get(W + "val")
        if current == tag:
            continue
        if tag in existing:
            raise ValueError(f"控件标签 {tag} 已用于其他单元格")
        if root.xpath("count(.//w:sdtPr/w:tag[@w:val=$value])", namespaces=NS, value=current) != 2:
            raise ValueError(f"原控件 {current} 不恰好出现两次，拒绝改动模板")
        nodes[0].set(W + "val", tag)
        existing.add(tag)
        changed = True
    if not changed:
        return
    parts[name] = (parts[name][0], etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True,
    ))
    write_docx_parts_atomic(parts, path)


def _cell_control_tag(path: Path, table_index: int, row: int, column: int) -> str:
    table = Document(path).tables[table_index - 1]
    if row < 1 or row > len(table.rows):
        raise ValueError(f"字段内容控件位置 {row} 行 {column} 列超出 Word 表格范围")
    cells = table.rows[row - 1]._tr.xpath("./w:tc")
    if column < 1 or column > len(cells):
        raise ValueError(f"字段内容控件位置 {row} 行 {column} 列超出 Word 表格范围")
    tags = cells[column - 1].xpath(".//w:sdtPr/w:tag/@w:val")
    if len(tags) != 1:
        raise ValueError(f"字段内容控件位置 {row} 行 {column} 列必须有且只有一个控件")
    return str(tags[0])


def _install_control_bindings(repository: Any, manifest: dict[str, Any],
                              template: Path, index: int, section: str) -> None:
    bindings = list(manifest.get("detailControls") or [])
    known = {item["controlTag"]: item for item in repository.list_mappings()
             if item.get("controlTag")}
    field_codes = {item["code"] for item in manifest["fields"]}
    for binding in bindings:
        field = str(binding.get("fieldCode") or "")
        if field not in field_codes or not repository.database.get_lims_field(field):
            raise ValueError(f"表格内容控件的标准字段 {field} 不存在")
        row, column = int(binding["row"]), int(binding["column"])
        tag = _cell_control_tag(template, index, row, column)
        if binding.get("tag") and tag != binding["tag"]:
            raise ValueError(f"表格内容控件 {row} 行 {column} 列的标签不符合配置")
        existing = known.get(tag)
        if existing:
            if existing.get("standardFieldCode") != field:
                raise ValueError(f"控件 {tag} 已绑定其他标准字段")
            continue
        repository.create_mapping({
            "locationId": f"word.content_control.{tag}", "sectionCode": section,
            "tableNo": manifest["tableNo"], "standardFieldCode": field,
            "fieldCode": f"report.segment.{manifest['tableNo']}.{row}.{column}",
            "sourceType": "SYSTEM", "repeatType": "ROW", "controlTag": tag,
            "enabled": True,
        })


def apply_segment_table(repository: Any, manifest: dict[str, Any]) -> int:
    active = repository.active_workspace()
    if not active or active["versionStatus"] != "DRAFT":
        raise ValueError("只能修改当前活动的草稿模板")
    table_no = str(manifest.get("tableNo") or "")
    header = str(manifest.get("tableHeader") or "")
    group = str(manifest.get("groupCode") or "")
    group_key = str(manifest.get("groupKey") or "")
    if not table_no or not header or not group or not group_key or not isinstance(manifest.get("tableLayout"), dict):
        raise ValueError("表格配置缺少表号、表头、编组身份字段或行片段版式")
    index, tags = target_table(Path(active["templateFile"]), header)
    mappings = [item for item in repository.list_mappings() if item.get("controlTag") in tags]
    if not mappings:
        raise ValueError(f"表格 {table_no} 尚未绑定任何系统字段")
    fields = {item["code"] for item in manifest["fields"]}
    rebind = manifest.get("rebindFields") or {}
    if not isinstance(rebind, dict) or any(
        old not in manifest.get("retireFields", []) or new not in fields
        for old, new in rebind.items()
    ):
        raise ValueError("表格字段重绑配置必须指向清单内的新字段")
    if any(item.get("standardFieldCode") not in fields and
           item.get("standardFieldCode") != manifest.get("conclusionField") and
           item.get("standardFieldCode") not in rebind for item in mappings):
        raise ValueError(f"表格 {table_no} 存在不属于目标编组的字段绑定")
    for replacement in set(rebind.values()):
        if not repository.database.get_lims_field(replacement):
            raise ValueError(f"重绑目标字段 {replacement} 不存在")
    for item in mappings:
        replacement = rebind.get(item.get("standardFieldCode"))
        if replacement:
            repository.update_mapping(item["id"], {"standardFieldCode": replacement})
    rules = [item for item in repository.list_table_rules() if item["tableNo"] == table_no]
    if len(rules) != 1:
        raise ValueError(f"表格 {table_no} 必须且只能存在一条表格规则")
    controls = manifest.get("uniqueControls") or []
    for item in controls:
        field = str(item.get("fieldCode") or "")
        if field not in fields:
            raise ValueError(f"唯一控件绑定的字段 {field} 不属于目标编组")
        if not repository.database.get_lims_field(field):
            raise ValueError(f"唯一控件的标准字段 {field} 不存在")
    _unique_tags(Path(active["templateFile"]), index, controls)
    _install_control_bindings(repository, manifest, Path(active["templateFile"]),
                              index, rules[0]["sectionCode"])
    for item in mappings:
        if item.get("repeatType") != "ROW" or item.get("tableNo") != table_no:
            repository.update_mapping(item["id"], {"repeatType": "ROW", "tableNo": table_no})
    repository.upsert_table_rule({**rules[0], "mode": "TABLE_REPEAT",
                                  "innerMode": "SEGMENT_REPEAT", "groupKey": group_key,
                                  "physicalTableIndex": index,
                                  "matrixLayout": manifest["tableLayout"]})
    repository.save_active_workspace()
    if controls:
        repository.set_version_document_key(active["versionId"], None)
    return index
