"""方案表格图片列使用与文本列相同的可配置定位和父记录分组。"""
import base64
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from backend.app.services.protocol_extractor import extract_protocol


PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='
)


@pytest.mark.parametrize('legacy', [False, True])
def test_configured_table_image_column_with_nested_records(tmp_path, legacy):
    doc = Document()
    doc.add_heading('杂质信息', 1)
    table = doc.add_table(rows=1, cols=3)
    for cell, title in zip(table.rows[0].cells, ('杂质名称', '结构式', '限度')):
        cell.text = title
    for name in ('杂质甲', '杂质乙'):
        row = table.add_row()
        row.cells[0].text = name
        row.cells[2].text = '0.1%'
        run = row.cells[1].paragraphs[0].add_run()
        run.add_picture(BytesIO(PNG))
        if legacy:
            blip = next(run._r.iter(qn('a:blip')))
            relationship = blip.get(qn('r:embed'))
            for child in list(run._r):
                run._r.remove(child)
            pict = etree.SubElement(run._r, qn('w:pict'))
            etree.SubElement(pict, '{urn:schemas-microsoft-com:vml}imagedata', {qn('r:id'): relationship})
    path = tmp_path / 'scheme.docx'
    doc.save(path)

    fields = [
        {'fieldCode': 'impurity.name', 'legacyJsonPath': '$.impurity[*].name',
         'collectionCode': 'impurity', 'dataType': 'string', 'enabled': True},
        {'fieldCode': 'impurity.image', 'legacyJsonPath': '$.impurity[*].injections[*].image',
         'collectionCode': 'impurity', 'dataType': 'image', 'enabled': True},
    ]
    mapping = {'sourceType': 'PROTOCOL', 'sectionPattern': '^杂质信息$',
               'headerPattern': '^杂质名称\\t结构式\\t限度$',
               'columnMappings': [
                   {'fieldCode': 'impurity.name', 'columnPattern': '^杂质名称$'},
                   {'fieldCode': 'impurity.image', 'columnPattern': '^结构式$'},
               ], 'rowExpansion': {'valueScope': 'PER_PARENT_ROW',
                                   'levelKey': 'injections', 'parentFieldCode': 'impurity.name'}}
    group = {'groupCode': 'impurity', 'cardinality': 'MANY', 'enabled': True,
             'sourceMappings': [mapping], 'fields': fields,
             'levels': [{'levelKey': 'injections', 'kind': 'ARRAY'}]}
    rules = [{'fieldCode': item['fieldCode'], 'name': '方案表格', 'sourceType': 'PROTOCOL',
              'transform': 'TRIM', 'enabled': True,
              'config': {'mode': 'TABLE_ROWS', 'groupCode': 'impurity', 'required': True}}
             for item in fields]

    result = extract_protocol(path, 'doc-1', fields, rules, [group], 4 * 1024 * 1024)

    assert [row['name'] for row in result['impurity']] == ['杂质甲', '杂质乙']
    assert all(base64.b64decode(row['injections'][0]['image'].split(',', 1)[1]) == PNG
               for row in result['impurity'])
    assert result['_meta']['fields']['impurity.image']['source']['locations'][0]['quote'] == '[图片]'


def test_image_rule_rejects_non_image_extraction_mode():
    from backend.app.services.protocol_rules import validate_protocol_rule

    field = {'fieldCode': 'impurity.image', 'legacyJsonPath': '$.impurity[*].image',
             'dataType': 'image'}
    with pytest.raises(ValueError, match='图片字段必须使用表格明细'):
        validate_protocol_rule(field, {'mode': 'RAW_BLOCK', 'sectionPattern': '杂质信息'}, [])
