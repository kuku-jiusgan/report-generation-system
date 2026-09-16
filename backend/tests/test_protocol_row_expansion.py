from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from docx import Document
from fastapi import APIRouter
import pytest

from backend.app.admin_routes.protocol_rules import register_protocol_rule_routes
from backend.app.services.protocol_extractor import extract_protocol
from backend.app.services.protocol_row_expansion import write_protocol_value
from backend.app.services.protocol_rules import validate_protocol_conflicts
from backend.app.services.protocol_report_source import replace_protocol_result


def configuration():
    fields = [{'fieldCode': code, 'label': key, 'collectionCode': 'summary',
               'legacyJsonPath': '$.summary[*].' + path, 'dataType': 'richText', 'enabled': True}
              for code, key, path in [('summary.item', '项目', 'item'),
                  ('summary.criteria', '标准', 'details[*].criteria')]]
    mapping = {'sourceType': 'PROTOCOL', 'sectionPattern': '汇总',
               'headerPattern': '^项目\t标准$', 'columnMappings': [
                   {'fieldCode': 'summary.item', 'columnPattern': '^项目$'},
                   {'fieldCode': 'summary.criteria', 'columnPattern': '^标准$'}],
               'rowExpansion': {'valueScope': 'PER_PARENT_ROW', 'levelKey': 'details',
                   'parentFieldCode': 'summary.item'}}
    groups = [{'groupCode': 'summary', 'enabled': True, 'cardinality': 'MANY',
               'fields': fields, 'levels': [{'levelKey': 'details', 'kind': 'ARRAY'}],
               'sourceMappings': [mapping]}]
    rules = [{'fieldCode': field['fieldCode'], 'sourceType': 'PROTOCOL', 'enabled': True,
              'transform': 'TRIM', 'config': {'mode': 'TABLE_ROWS', 'groupCode': 'summary'}}
             for field in fields]
    return fields, rules, groups


def document(tmp_path, rows=None):
    doc = Document()
    doc.add_heading('汇总', 1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = '项目', '标准'
    for item, criteria in rows or [('系统适用性', ' RSD≤20%\n保留时间稳定 '), ('准确度', ' 回收率80%～120% ')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = item, criteria
    path = tmp_path / '方案.docx'
    doc.save(path)
    return path


def extract(path, config=None):
    fields, rules, groups = config or configuration()
    return extract_protocol(path, 'doc', fields, rules, groups, 4 * 1024 * 1024)


def test_criteria_belong_to_the_project_on_the_same_source_row(tmp_path):
    result = extract(document(tmp_path))
    assert result['summary'] == [
        {'item': '系统适用性', 'details': [{'criteria': 'RSD≤20%\n保留时间稳定'}]},
        {'item': '准确度', 'details': [{'criteria': '回收率80%～120%'}]},
    ]
    criteria = result['_meta']['fields']['summary.criteria']
    assert [location['parentValue'] for location in criteria['source']['locations']] == ['系统适用性', '准确度']
    assert [location['row'] for location in criteria['source']['locations']] == [2, 3]
    data = {'source_payloads': {'EXCEL': {'summary': [{'keep': True}]}}}
    replace_protocol_result(data, result, 'hash')
    assert data['source_payloads']['EXCEL'] == {'summary': [{'keep': True}]}


def test_expansion_does_not_depend_on_field_rule_order(tmp_path):
    fields, rules, groups = configuration()
    rules.reverse()
    assert extract(document(tmp_path), (fields, rules, groups))['summary'] == extract(document(tmp_path))['summary']


def test_noncontiguous_rows_group_by_project_with_uneven_detail_counts(tmp_path):
    result = extract(document(tmp_path, [('项目甲', 'A1'), ('项目乙', 'B1'), (' 项目甲 ', 'A2')]))
    assert result['summary'] == [
        {'item': '项目甲', 'details': [{'criteria': 'A1'}, {'criteria': 'A2'}]},
        {'item': '项目乙', 'details': [{'criteria': 'B1'}]},
    ]


def test_empty_project_key_is_an_error(tmp_path):
    with pytest.raises(ValueError, match='分组键不能为空'):
        extract(document(tmp_path, [('', '标准')]))


def test_missing_parent_column_mapping_is_an_error():
    fields, rules, groups = configuration()
    groups[0]['sourceMappings'][0]['columnMappings'].pop(0)
    with pytest.raises(ValueError, match='分组键必须配置主表列映射'):
        validate_protocol_conflicts(fields, rules, groups)


def test_nested_configuration_requires_explicit_scope_and_array_level():
    fields, rules, groups = configuration()
    groups[0]['sourceMappings'][0]['rowExpansion'].pop('valueScope')
    with pytest.raises(ValueError, match='父记录键分组'):
        validate_protocol_conflicts(fields, rules, groups)
    fields, rules, groups = configuration()
    groups[0]['levels'][0]['kind'] = 'OBJECT'
    with pytest.raises(ValueError, match='数组层不存在'):
        validate_protocol_conflicts(fields, rules, groups)


def test_nested_write_keeps_uneven_boundaries_and_rejects_mismatch_atomically():
    payload = {}
    write_protocol_value(payload, '$.summary[*].details[*].identity', [['甲'], ['乙', '丙']])
    write_protocol_value(payload, '$.summary[*].details[*].criteria', [['A'], ['B', 'C']])
    assert payload['summary'][1]['details'][1] == {'identity': '丙', 'criteria': 'C'}
    before = deepcopy(payload)
    with pytest.raises(ValueError, match='数量不一致'):
        write_protocol_value(payload, '$.summary[*].details[*].criteria', [['A'], ['B']])
    assert payload == before
    with pytest.raises(ValueError, match='父子分组值'):
        write_protocol_value(payload, '$.summary[*].details[*].criteria', ['A', 'B', 'C'])


def test_rule_editor_preview_uses_unsaved_mapping_and_shows_parent_records(tmp_path):
    fields, rules, groups = configuration()
    before = deepcopy(groups)
    path = document(tmp_path)
    doc = Document(path)
    doc.paragraphs[0].text = '可编辑章节'
    doc.tables[0].cell(0, 0).text = '项目名称'
    doc.tables[0].cell(0, 1).text = '判定要求'
    doc.save(path)
    repository = MagicMock()
    repository.database.get_lims_field.side_effect = lambda code: next(field for field in fields if field['fieldCode'] == code)
    repository.database.list_system_field_rules.return_value = [rules[0]]
    repository.database.get_source.return_value = {'id': 'doc', 'source_type': 'PROTOCOL', 'stored_name': path.name}
    router = APIRouter()
    register_protocol_rule_routes(router, repository, SimpleNamespace(uploads_dir=tmp_path, max_upload_mb=4))
    preview = next(route.endpoint for route in router.routes if route.path.endswith('/preview'))
    mapping = deepcopy(groups[0]['sourceMappings'][0])
    mapping.update(sectionPattern='可编辑章节', headerPattern='^项目名称\t判定要求$', rowPattern='准确度')
    mapping['columnMappings'][0]['columnPattern'] = '^项目名称$'
    mapping['columnMappings'][1]['columnPattern'] = '^判定要求$'
    with patch('backend.app.admin_routes.protocol_rules.list_system_field_groups', return_value=groups):
        result = preview({'documentId': 'doc', 'fieldCode': 'summary.criteria',
                          'config': rules[1]['config'], 'sourceMappings': [mapping]})
    assert result['groups']['summary'] == [{'item': '准确度', 'details': [{'criteria': '回收率80%～120%'}]}]
    assert result['fields']['summary.item']['value'] == ['准确度']
    assert groups == before
    repository.database.update_source_payload.assert_not_called()
    repository.database.update_report.assert_not_called()
    repository.database.save_system_field_rule.assert_not_called()


def test_preview_rejects_parent_rule_from_another_source():
    from backend.app.admin_routes.protocol_rules import _preview_inputs
    fields, rules, groups = configuration()
    database = MagicMock()
    database.get_lims_field.return_value = fields[0]
    database.list_system_field_rules.return_value = [{**rules[0], 'sourceType': 'LIMS'}]
    with pytest.raises(ValueError, match='父记录分组键必须启用同一编组'):
        _preview_inputs(database, fields[1], rules[1], groups, {})
