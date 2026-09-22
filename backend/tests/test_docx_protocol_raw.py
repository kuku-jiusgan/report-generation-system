"""方案连续原文块按字段内容控件复制，范围内可同时包含段落和表格。"""

import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from backend.app.services.mapped_docx_generator import build_mapped_docx
from backend.app.services.docx_protocol_raw import _copy_blocks, _source_style_closure

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W_NS}


def _target(path: Path) -> None:
    body = (
        f'<w:p><w:r><w:t>9. 计算公式及修约原则</w:t></w:r></w:p>'
        f'<w:sdt><w:sdtPr><w:tag w:val="cc.raw"/></w:sdtPr>'
        f'<w:sdtContent><w:p><w:r><w:t>空行</w:t></w:r></w:p></w:sdtContent></w:sdt>'
        f'<w:sectPr/>'
    )
    document = f'<w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>{body}</w:body></w:document>'
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/styles.xml", f'<w:styles xmlns:w="{W_NS}"/>')


def _source(path: Path) -> None:
    document = Document()
    document.add_heading("9. 计算公式及修约原则", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("含格式的公式：")
    run = paragraph.add_run("x")
    run.bold = True
    document.add_paragraph("结果保留两位小数。")
    document.add_heading("10. 偏差", level=1)
    document.add_paragraph("不应复制。")
    document.save(path)


def test_field_bound_raw_section_replaces_empty_control(tmp_path: Path) -> None:
    target, source, output = (tmp_path / name for name in ("template.docx", "protocol.docx", "report.docx"))
    _target(target)
    _source(source)
    mapping = {"fieldCode": "report.s9.formula", "standardFieldCode": "method.formula",
               "controlTag": "cc.raw", "enabled": True}
    rule = {"fieldCode": "method.formula", "sourceType": "PROTOCOL", "enabled": True,
            "config": {"mode": "RAW_BLOCK",
                       "sectionPattern": "^9\\. 计算公式及修约原则$", "includeStart": False}}
    build_mapped_docx(target, output, [mapping], {"method": {"formula": "预览文本"}}, {},
                      protocol_document=source, protocol_rules=[rule])
    with zipfile.ZipFile(output) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
    assert document.xpath(".//w:sdtContent//w:t/text()", namespaces=NS) == [
        "含格式的公式：", "x", "结果保留两位小数。",
    ]
    assert document.xpath(".//w:sdtContent//w:rPr/w:b", namespaces=NS)
    assert "空行" not in "".join(document.xpath(".//w:t/text()", namespaces=NS))


def test_continuous_raw_block_keeps_paragraphs_and_table(tmp_path: Path) -> None:
    target, source, output = (tmp_path / name for name in ("template.docx", "protocol.docx", "report.docx"))
    _target(target)
    document = Document()
    document.add_heading("6. 分析方法", level=1)
    document.add_paragraph("表61 方法参数列表")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "参数"
    document.add_paragraph("备注：不属于方法参数表。")
    document.add_heading("7. 验证内容", level=1)
    document.save(source)
    mapping = {"fieldCode": "report.s6.method", "standardFieldCode": "method.field_001",
               "controlTag": "cc.raw", "enabled": True}
    rule = {"fieldCode": "method.field_001", "sourceType": "PROTOCOL", "enabled": True,
            "config": {"mode": "RAW_BLOCK", "sectionPattern": "^6\\. 分析方法$",
                       "includeStart": False}}
    build_mapped_docx(target, output, [mapping], {}, {}, protocol_document=source,
                      protocol_rules=[rule])
    with zipfile.ZipFile(output) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
    content = document.xpath(".//w:sdtContent", namespaces=NS)[0]
    assert [child.tag.rsplit("}", 1)[-1] for child in content] == ["p", "tbl", "p"]
    assert content.xpath(".//w:tbl//w:t/text()", namespaces=NS) == ["项目", "参数"]
    assert content.xpath("./w:p[last()]//w:t/text()", namespaces=NS) == ["备注：不属于方法参数表。"]


def test_style_dependencies_are_collected_and_source_bookmarks_are_removed() -> None:
    styles = etree.fromstring(
        f'<w:styles xmlns:w="{W_NS}">'
        '<w:style w:type="paragraph" w:styleId="base"/>'
        '<w:style w:type="paragraph" w:styleId="child"><w:basedOn w:val="base"/></w:style>'
        '</w:styles>'
    )
    assert set(_source_style_closure(styles, {"child"})) == {"base", "child"}
    paragraph = etree.fromstring(
        f'<w:p xmlns:w="{W_NS}"><w:bookmarkStart w:id="1" w:name="source"/>'
        '<w:r><w:t>正文</w:t></w:r><w:bookmarkEnd w:id="1"/></w:p>'
    )
    copied = _copy_blocks([paragraph])[0]
    assert copied.xpath(".//w:t/text()", namespaces=NS) == ["正文"]
    assert not copied.xpath(".//w:bookmarkStart | .//w:bookmarkEnd", namespaces=NS)
