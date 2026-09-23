from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from fastapi import APIRouter

from backend.app.admin_routes.protocol_rules import register_protocol_rule_routes
from backend.app.services.docx_field_values import source_mapping_value
from backend.app.services.protocol_extractor import extract_protocol
from backend.app.services.protocol_report_source import refresh_protocol_source, replace_protocol_result
from backend.app.services.protocol_rules import validate_protocol_conflicts
from backend.app.services.system_field_resolver import resolve_system_fields
from backend.app.services.system_field_groups import _validate_group_contract

MAX_BYTES = 4 * 1024 * 1024


def field(code='project.code', path='$.project.code', group='project'):
    return {'fieldCode': code, 'label': code, 'legacyJsonPath': path,
            'collectionCode': group, 'dataType': 'string', 'enabled': True}


def rule(code='project.code', mode='LABEL', transform='TRIM', **config):
    return {'id': 41, 'fieldCode': code, 'name': '方案规则', 'sourceType': 'PROTOCOL',
            'transform': transform, 'enabled': True, 'config': {'mode': mode, 'required': True,
            **({'sectionPattern': '概述', 'labelPattern': '方案编号'} if mode == 'LABEL' else {}), **config}}


def scheme(tmp_path, duplicate=False):
    document = Document()
    document.add_heading('概述', 1)
    paragraph = document.add_paragraph()
    paragraph.add_run('方案')
    paragraph.add_run('编号： P-2026')
    if duplicate:
        document.add_paragraph('方案编号：P-2027')
    document.add_heading('目的', 2)
    document.add_paragraph('第一段原文。')
    document.add_paragraph('第二段原文。')
    document.add_heading('材料', 1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = '名称'
    table.rows[0].cells[1].text = '型号'
    for name, model in [('仪器甲', 'A'), ('仪器乙', 'B')]:
        row = table.add_row()
        row.cells[0].text = name
        row.cells[1].text = model
    document.add_heading('结束', 1)
    document.add_paragraph('不应提取。')
    path = tmp_path / '方案.docx'
    document.save(path)
    return path


def extract(path, fields=None, rules=None, groups=None, strict=True):
    return extract_protocol(path, 'doc-1', fields or [field()], rules or [rule()], groups or [], MAX_BYTES, strict)


def detail_config():
    fields = [field('instruments.name', '$.instruments[*].name', 'instruments'),
              field('instruments.model', '$.instruments[*].model', 'instruments')]
    group = {'groupCode': 'instruments', 'cardinality': 'MANY', 'enabled': True,
             'sourceMappings': [{'sourceType': 'PROTOCOL', 'sectionPattern': '材料',
               'headerPattern': '^名称\\t型号$', 'columnMappings': [
                   {'fieldCode': 'instruments.name', 'columnPattern': '^名称$'},
                   {'fieldCode': 'instruments.model', 'columnPattern': '^型号$'}]}]}
    rules = [rule(item['fieldCode'], 'TABLE_ROWS', groupCode='instruments') for item in fields]
    return fields, rules, [group]


def test_label_across_runs_and_provenance(tmp_path):
    result = extract(scheme(tmp_path))
    assert result['project']['code'] == 'P-2026'
    source = result['_meta']['fields']['project.code']['source']
    assert source['type'] == 'PROTOCOL'
    assert source['document_id'] == 'doc-1'
    assert source['locations'][0]['paragraph'] == 2
    assert source['locations'][0]['quote'] == '方案编号： P-2026'


def test_regex_replace_transforms_scalar_protocol_result(tmp_path):
    configured = rule(transform='REGEX_REPLACE', replacePattern='^P-', replaceWith='R-')

    result = extract(scheme(tmp_path), rules=[configured])

    assert result['project']['code'] == 'R-2026'


def test_document_code_from_header_table_uses_field_rule(tmp_path):
    document = Document()
    table = document.sections[0].header.add_table(rows=1, cols=2, width=1)
    table.cell(0, 0).text = '文件编号'
    table.cell(0, 1).text = 'ZBYY/MV-R XM2026234-02'
    path = tmp_path / '页眉编号方案.docx'
    document.save(path)

    result = extract(path, fields=[field('document.code', '$.document.code', 'document')],
        rules=[rule('document.code', 'HEADER_TABLE_CELL', labelPattern='^文件编号$')])

    assert result['document']['code'] == 'ZBYY/MV-R XM2026234-02'
    location = result['_meta']['fields']['document.code']['source']['locations'][0]
    assert location == {'section': '页眉', 'table': 1, 'row': 1, 'column': 2,
                        'quote': 'ZBYY/MV-R XM2026234-02'}


def test_header_title_is_selected_and_replaced_by_field_rule(tmp_path):
    document = Document()
    table = document.sections[0].header.add_table(rows=3, cols=4, width=1)
    for row, marker in zip(table.rows, ('文件编号', '版 本 号', '页    码')):
        row.cells[0].text = '文件标题'
        row.cells[1].text = '加替沙星原料药中N-亚硝基加替沙星分析方法\n验证方案'
        row.cells[2].text = marker
    path = tmp_path / '页眉标题方案.docx'
    document.save(path)
    title_rule = rule('header.field_001', 'HEADER_TABLE_CELL', labelPattern='^文件标题$',
        transform='REGEX_REPLACE', rowPattern='文件编号',
        replacePattern='验证方案$', replaceWith='验证报告')

    result = extract(path, fields=[field('header.field_001', '$.document.field_001', 'document')],
                     rules=[title_rule])

    assert result['document']['field_001'] == '加替沙星原料药中N-亚硝基加替沙星分析方法\n验证报告'
    assert result['_meta']['fields']['header.field_001']['source']['locations'][0]['quote'].endswith('验证方案')


def test_header_replacement_must_match_once(tmp_path):
    document = Document()
    table = document.sections[0].header.add_table(rows=1, cols=2, width=1)
    table.cell(0, 0).text = '文件标题'
    table.cell(0, 1).text = '没有待替换内容'
    path = tmp_path / '页眉标题不匹配方案.docx'
    document.save(path)
    title_rule = rule('header.field_001', 'HEADER_TABLE_CELL', labelPattern='^文件标题$',
        transform='REGEX_REPLACE', replacePattern='验证方案$', replaceWith='验证报告')

    with pytest.raises(ValueError, match='结果替换正则：未匹配'):
        extract(path, fields=[field('header.field_001', '$.document.field_001', 'document')],
                rules=[title_rule])


def test_section_has_exact_boundary_and_paragraphs(tmp_path):
    result = extract(scheme(tmp_path), rules=[rule(mode='SECTION', sectionPattern='概述 / 目的')])
    assert result['project']['code'] == '第一段原文。\n第二段原文。'


def test_raw_block_extracts_plain_preview_for_word_copy(tmp_path):
    result = extract(scheme(tmp_path), rules=[rule(
        mode='RAW_BLOCK', sectionPattern='概述 / 目的', includeStart=False,
    )])
    assert result['project']['code'] == '第一段原文。\n第二段原文。'
    assert result['_meta']['fields']['project.code']['source']['locations'][0]['section'] == '概述 / 目的'


def test_validation_methods_and_criteria_from_protocol_sections(tmp_path):
    projects = [
        ('validation.field_001', 'validation.field_002', '系统适用性'),
        ('validation.field_004', 'validation.field_005', '专属性'),
        ('validation.field_007', 'validation.field_008', '检测限与定量限'),
        ('validation.field_010', 'validation.field_011', '线性与范围'),
        ('validation.field_013', 'validation.field_014', '重复性'),
    ]
    document = Document()
    document.add_heading('验证内容', 1)
    for _, _, project in projects:
        document.add_heading(project, 2)
        document.add_heading('试验方法', 3)
        document.add_paragraph(f'{project}试验方法正文。')
        document.add_heading('接受标准', 3)
        document.add_paragraph(f'{project}接受标准正文。')
    path = tmp_path / '方案.docx'
    document.save(path)
    codes = [code for method, criteria, _ in projects for code in (method, criteria)]
    fields = [field(code, f'$.custom.{code.rsplit(".", 1)[-1]}', 'custom') for code in codes]
    rules = [rule(code, 'SECTION', sectionPattern=rf'^验证内容 / {project} / {kind}$')
             for method, criteria, project in projects
             for code, kind in ((method, '试验方法'), (criteria, '接受标准'))]

    result = extract(path, fields, rules)

    for method, criteria, project in projects:
        assert result['custom'][method.rsplit('.', 1)[-1]] == f'{project}试验方法正文。'
        assert result['custom'][criteria.rsplit('.', 1)[-1]] == f'{project}接受标准正文。'


def test_duplicate_label_required_and_optional(tmp_path):
    path = scheme(tmp_path, duplicate=True)
    with pytest.raises(ValueError, match='匹配到 2 处'):
        extract(path)
    result = extract(path, rules=[rule(required=False)])
    assert 'project' not in result
    assert result['_meta']['fields']['project.code']['status'] == 'ERROR'
    assert result['_meta']['warnings']
    assert not result['_meta']['errors']


def test_missing_document_policy():
    with pytest.raises(ValueError, match='未上传'):
        extract(None)
    assert extract(None, rules=[rule(required=False)])['_meta']['warnings']
    assert extract_protocol(None, '', [], [], [], MAX_BYTES)['_meta']['fields'] == {}


def test_table_rows_share_record_order(tmp_path):
    fields, rules, groups = detail_config()
    result = extract(scheme(tmp_path), fields, rules, groups)
    assert result['instruments'] == [{'name': '仪器甲', 'model': 'A'}, {'name': '仪器乙', 'model': 'B'}]
    sources = result['_meta']['fields']
    assert [item['row'] for item in sources['instruments.name']['source']['locations']] == [2, 3]
    assert [item['row'] for item in sources['instruments.model']['source']['locations']] == [2, 3]


def test_regex_replace_transforms_each_protocol_table_row(tmp_path):
    fields, rules, groups = detail_config()
    rules[0].update({'transform': 'REGEX_REPLACE'})
    rules[0]['config'].update({'replacePattern': '^仪器', 'replaceWith': '设备'})

    result = extract(scheme(tmp_path), fields, rules, groups)

    assert [item['name'] for item in result['instruments']] == ['设备甲', '设备乙']


def test_table_filter_and_single_value(tmp_path):
    result = extract(scheme(tmp_path), rules=[rule(mode='TABLE_COLUMN', sectionPattern='材料',
        headerPattern='^名称\\t型号$', columnPattern='^型号$', rowPattern='仪器乙')])
    assert result['project']['code'] == 'B'
    with pytest.raises(ValueError, match='单值行.*2 处'):
        extract(scheme(tmp_path), rules=[rule(mode='TABLE_COLUMN', sectionPattern='材料',
            headerPattern='^名称\\t型号$', columnPattern='^型号$')])
    result = extract(scheme(tmp_path), rules=[rule(mode='TABLE_CELL', sectionPattern='材料',
        headerPattern='^名称\\t型号$', labelPattern='仪器甲')])
    assert result['project']['code'] == 'A'


def test_empty_cell_does_not_shift_rows(tmp_path):
    path = scheme(tmp_path)
    document = Document(path)
    document.tables[0].rows[1].cells[1].text = ''
    document.save(path)
    fields, rules, groups = detail_config()
    rules[1]['config']['required'] = False
    result = extract(path, fields, rules, groups)
    assert result['instruments'] == [{'name': '仪器甲'}, {'name': '仪器乙'}]
    assert result['_meta']['fields']['instruments.model']['status'] == 'ERROR'
    assert '第 2 行' in result['_meta']['warnings'][0]


def test_vertical_merged_cells_keep_alignment(tmp_path):
    path = scheme(tmp_path)
    document = Document(path)
    document.tables[0].cell(1, 0).merge(document.tables[0].cell(2, 0)).text = '共享名称'
    document.save(path)
    fields, rules, groups = detail_config()
    result = extract(path, fields, rules, groups)
    assert result['instruments'] == [{'name': '共享名称', 'model': 'A'}, {'name': '共享名称', 'model': 'B'}]


def test_unstyled_section_requires_explicit_end(tmp_path):
    document = Document()
    document.add_paragraph('概述')
    document.add_paragraph('方案编号：ABC')
    document.add_paragraph('结束')
    path = tmp_path / '方案.docx'
    document.save(path)
    with pytest.raises(ValueError, match='没有标题层级'):
        extract(path, rules=[rule(required=False)])
    assert extract(path, rules=[rule(endPattern='^结束$')])['project']['code'] == 'ABC'


def test_rule_conflicts_and_invalid_regex_are_explicit():
    with pytest.raises(ValueError, match='不能同时启用'):
        validate_protocol_conflicts([field()], [rule(), {'fieldCode': 'project.code', 'sourceType': 'LIMS'}], [])
    with pytest.raises(ValueError, match='正则无效'):
        validate_protocol_conflicts([field()], [rule(labelPattern='[')], [])
    with pytest.raises(ValueError, match='必须配置替换正则'):
        validate_protocol_conflicts([field()], [rule(transform='REGEX_REPLACE')], [])
    with pytest.raises(ValueError, match='明细展开'):
        fields, rules, groups = detail_config()
        fields[0]['legacyJsonPath'] = '$.instruments[*].details[*].name'
        validate_protocol_conflicts(fields, rules, groups)


def test_source_namespace_and_calculation(tmp_path):
    result = extract(scheme(tmp_path))
    data = {'source_payloads': {'EXCEL': {'project': {'other': 'keep'}}}}
    replace_protocol_result(data, result, 'hash')
    fields = [field(), field('project.description', '$.project.description')]
    rules = [rule(), {'fieldCode': 'project.description', 'sourceType': 'CALCULATED',
        'config': {'dependencies': ['project.code'], 'textTemplate': '方案 {project.code}'}}]
    resolve_system_fields(fields, rules, data['source_payloads']['EXCEL'], data)
    assert 'code' not in data['source_payloads']['EXCEL']['project']
    assert data['source_payloads']['EXCEL']['project']['description'] == '方案 P-2026'
    assert data['field_sources']['project.code']['locations']
    assert source_mapping_value({'standardFieldCode': 'project.code', 'sourcePath': '$.project.code'},
                                data['source_payloads']['EXCEL'], data) == 'P-2026'


def test_reextract_removes_stale_values_and_warning(tmp_path):
    data = {'source_payloads': {'LIMS': {'project': {'other': 'keep'}}}}
    replace_protocol_result(data, extract(scheme(tmp_path)), 'old')
    replacement = extract(scheme(tmp_path), rules=[rule(labelPattern='找不到', required=False)])
    replace_protocol_result(data, replacement, 'new')
    assert 'project.code' not in data['original_values']
    assert 'project' not in data['source_payloads']['PROTOCOL']
    assert data['source_payloads']['LIMS'] == {'project': {'other': 'keep'}}
    replace_protocol_result(data, extract_protocol(None, '', [], [], [], MAX_BYTES), '')
    assert 'project.code' not in data['field_sources']
    assert not data['warnings']


def test_runtime_missing_document_is_kept_as_warning(tmp_path):
    database = MagicMock()
    database.list_lims_fields.return_value = [field()]
    database.list_system_field_rules.return_value = [rule()]
    settings = SimpleNamespace(uploads_dir=tmp_path, max_upload_mb=4)
    data = {'source_payloads': {'EXCEL': {'keep': True}}}
    with patch('backend.app.services.system_field_groups.list_system_field_groups', return_value=[]):
        refresh_protocol_source(database, settings, data)
    assert data['source_payloads']['EXCEL'] == {'keep': True}
    assert data['source_payloads']['PROTOCOL']['_meta']['warnings']
    assert '未上传' in data['source_payloads']['PROTOCOL']['_meta']['warnings'][0]


def test_runtime_locator_failure_is_kept_as_empty_value(tmp_path):
    database = MagicMock()
    database.list_lims_fields.return_value = [field()]
    database.list_system_field_rules.return_value = [rule(sectionPattern='不存在的章节')]
    source_path = scheme(tmp_path)
    database.get_source.return_value = {
        'id': 'doc-1', 'source_type': 'PROTOCOL', 'stored_name': source_path.name,
        'sha256': 'hash',
    }
    settings = SimpleNamespace(uploads_dir=tmp_path, max_upload_mb=4)
    data = {'source_payloads': {'PROTOCOL_DOCUMENT': {'id': 'doc-1'}, 'EXCEL': {'keep': True}}}
    with patch('backend.app.services.system_field_groups.list_system_field_groups', return_value=[]):
        refresh_protocol_source(database, settings, data)
    result = data['source_payloads']['PROTOCOL']['_meta']['fields']['project.code']
    assert result['status'] == 'ERROR'
    assert result.get('value') is None
    assert '未匹配' in result['message']
    assert data['source_payloads']['EXCEL'] == {'keep': True}


def test_preview_is_read_only_and_shows_required_failure(tmp_path):
    path = scheme(tmp_path, duplicate=True)
    repository = MagicMock()
    repository.database.get_lims_field.return_value = field()
    repository.database.get_source.return_value = {'id': 'doc-1', 'source_type': 'PROTOCOL', 'stored_name': path.name}
    router = APIRouter()
    register_protocol_rule_routes(router, repository, SimpleNamespace(uploads_dir=tmp_path, max_upload_mb=4))
    preview = next(route.endpoint for route in router.routes if route.path.endswith('/preview'))
    with patch('backend.app.admin_routes.protocol_rules.list_system_field_groups', return_value=[]):
        result = preview({'documentId': 'doc-1', 'fieldCode': 'project.code', 'config': rule()['config']})
    assert result['errors']
    repository.database.update_source_payload.assert_not_called()
    repository.database.update_report.assert_not_called()


def test_protocol_metadata_exposes_shared_result_transforms():
    repository = MagicMock()
    repository.database.list_sources.return_value = []
    router = APIRouter()
    register_protocol_rule_routes(router, repository, SimpleNamespace())
    metadata = next(route.endpoint for route in router.routes
                    if route.path.endswith('/protocol-rule-metadata'))()

    assert metadata['transforms'][-1] == {'value': 'REGEX_REPLACE', 'label': '正则替换'}
    assert metadata['transformGroups'][0]['fields'][0]['key'] == 'replacePattern'


def test_group_locator_validation():
    fields, _, groups = detail_config()
    group = groups[0]
    _validate_group_contract(group, fields)
    group['sourceMappings'][0]['rowPattern'] = '['
    with pytest.raises(ValueError, match='正则无效'):
        _validate_group_contract(group, fields)


def test_rule_save_replaces_current_source_without_conflict():
    from backend.app.admin_routes.rule_catalog import _validate_system_rule
    repository = MagicMock()
    repository.database.get_lims_field.return_value = field()
    repository.database.list_lims_fields.return_value = [field()]
    repository.database.list_system_field_rules.return_value = [
        {'id': 41, 'fieldCode': 'project.code', 'sourceType': 'LIMS', 'enabled': True}]
    with patch('backend.app.admin_routes.rule_catalog.list_system_field_groups', return_value=[]):
        validated = _validate_system_rule(repository, rule())
    assert validated['sourceType'] == 'PROTOCOL'
    with patch('backend.app.admin_routes.rule_catalog.list_system_field_groups', return_value=[]):
        from fastapi import HTTPException
        with pytest.raises(HTTPException, match='不能同时启用'):
            _validate_system_rule(repository, {**rule(), 'id': None})


def test_optional_missing_entire_table_has_empty_records(tmp_path):
    fields, rules, groups = detail_config()
    groups[0]['sourceMappings'][0]['headerPattern'] = '不存在的表头'
    for item in rules:
        item['config']['required'] = False
    result = extract(scheme(tmp_path), fields, rules, groups)
    assert result['instruments'] == []
    assert len(result['_meta']['warnings']) == 2


def test_runtime_reextracts_original_document_with_current_rule(tmp_path):
    path = scheme(tmp_path)
    database = MagicMock()
    database.list_lims_fields.return_value = [field()]
    database.list_system_field_rules.return_value = [rule()]
    database.get_source.return_value = {'id': 'doc-1', 'source_type': 'PROTOCOL', 'stored_name': path.name, 'sha256': 'hash'}
    settings = SimpleNamespace(uploads_dir=tmp_path, max_upload_mb=4)
    data = {'source_payloads': {'PROTOCOL_DOCUMENT': {'id': 'doc-1'}, 'EXCEL': {'keep': True}}}
    with patch('backend.app.services.system_field_groups.list_system_field_groups', return_value=[]):
        refresh_protocol_source(database, settings, data)
        assert data['original_values']['project.code'] == 'P-2026'
        database.list_system_field_rules.return_value = [rule(labelPattern='未配置标签', required=False)]
        refresh_protocol_source(database, settings, data)
        assert 'project.code' not in data['original_values']
        assert data['field_sources']['project.code']['message']
    assert data['source_payloads']['EXCEL'] == {'keep': True}


def test_unrelated_nested_table_does_not_block_selected_section(tmp_path):
    path = scheme(tmp_path)
    document = Document(path)
    document.add_heading('附表', 1)
    document.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
    document.save(path)
    assert extract(path)['project']['code'] == 'P-2026'
    with pytest.raises(ValueError, match='嵌套表格'):
        extract(path, rules=[rule(mode='SECTION', sectionPattern='附表', required=False)])


def test_acceptance_criteria_from_protocol_summary_table(tmp_path):
    document = Document()
    document.add_heading('结果汇总', 1)
    document.add_heading('验证结果汇总', 2)
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['验证项目', '接受标准', '验证结论']):
        cell.text = text
    expected = ['峰面积RSD≤20%\n保留时间RSD≤1%', '回收率为80%～120%']
    for item, criteria in zip(['系统适用性', '准确度'], expected):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = item, criteria
    path = tmp_path / '方案.docx'
    document.save(path)
    code = 'validationSummary.acceptanceCriteria'
    fields = [{**field(code, '$.validationSummary[*].acceptanceCriteria', 'validationSummary'),
               'dataType': 'richText'}]
    rules = [rule(code, 'TABLE_ROWS', groupCode='validationSummary')]
    groups = [{'groupCode': 'validationSummary', 'cardinality': 'MANY', 'enabled': True,
               'sourceMappings': [{'sourceType': 'PROTOCOL',
                   'sectionPattern': '结果汇总 / 验证结果汇总',
                   'headerPattern': '^验证项目\\t接受标准\\t验证结论$',
                   'columnMappings': [{'fieldCode': code, 'columnPattern': '^接受标准$'}]}]}]
    result = extract(path, fields, rules, groups)
    assert [row['acceptanceCriteria'] for row in result['validationSummary']] == expected
    assert [location['row'] for location in result['_meta']['fields'][code]['source']['locations']] == [2, 3]
