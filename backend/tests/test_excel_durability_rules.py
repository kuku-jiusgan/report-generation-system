from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook

from backend.app.services.excel_field_extractor import extract_excel_fields
from backend.app.services.excel_rule_defaults import (
    DURABILITY_FIELDS,
    EXCEL_FIELD_PATHS,
    _rule_config,
    ensure_excel_field_rules,
)


GROUP = "custom_1789628945793"
ROOT = Path(__file__).parents[2]


def _fields_and_rules() -> tuple[list[dict], list[dict]]:
    fields = []
    rules = []
    for index, code in enumerate(DURABILITY_FIELDS, 1):
        path = EXCEL_FIELD_PATHS[code]
        fields.append({
            "fieldCode": code,
            "groupCode": GROUP,
            "cardinality": "ONE" if code == "uncategorized.field_122" else "MANY",
            "legacyJsonPath": path,
        })
        rules.append({
            "id": index,
            "fieldCode": code,
            "sourceType": "EXCEL",
            "enabled": True,
            "priority": 50,
            "config": _rule_config(code, path),
        })
    return fields, rules


def _workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "耐用性"
    sheet.append(["杂质名称", "杂质 A", None])
    sheet.append(["溶液名称", "色谱柱 1", "色谱柱 2"])
    sheet.append(["空白溶液", "blank-1", "blank-2"])
    sheet.append(["供试品溶液", "sample-1", "sample-2"])
    sheet.append(["线性方程", "equation-1", "equation-2"])
    sheet.append(["线性系数", "coef-1", "coef-2"])
    sheet.append(["加标溶液测得浓度", "concentration-1", "concentration-2"])
    sheet.append(["浓度比值", "ratio", None])
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("B8:C8")
    workbook.save(path)


def test_durability_rules_split_one_impurity_into_two_injections(tmp_path: Path) -> None:
    workbook = tmp_path / "durability.xlsm"
    _workbook(workbook)
    fields, rules = _fields_and_rules()

    payload = extract_excel_fields(workbook, fields, rules)

    records = payload[GROUP]
    assert len(records) == 1
    assert records[0]["field_122"] == "杂质 A"
    assert records[0]["injections"] == [
        {
            "field_097": "色谱柱 1",
            "field_098": "blank-1",
            "field_099": "sample-1",
            "field_100": "equation-1",
            "field_101": "coef-1",
            "field_102": "concentration-1",
            "field_103": "ratio",
        },
        {
            "field_097": "色谱柱 2",
            "field_098": "blank-2",
            "field_099": "sample-2",
            "field_100": "equation-2",
            "field_101": "coef-2",
            "field_102": "concentration-2",
            "field_103": "ratio",
        },
    ]
    assert not payload["_meta"]["warnings"]


def test_durability_rule_can_be_edited_from_catalog_without_being_reset() -> None:
    database = Mock()
    field_code = "uncategorized.field_097"
    custom_config = {"sourcePath": EXCEL_FIELD_PATHS[field_code], "mode": "REPEAT_BLOCK",
                     "sheet": "耐用性", "rowStart": 12, "rowEnd": 12,
                     "repeatCount": 1, "valueMode": "CELL"}
    database.get_lims_field.side_effect = lambda code: (
        {"fieldCode": field_code, "groupCode": GROUP,
         "legacyJsonPath": EXCEL_FIELD_PATHS[field_code]}
        if code == field_code else None
    )
    database.list_system_field_rules.return_value = [{
        "id": 1, "fieldCode": field_code, "sourceType": "EXCEL", "enabled": True,
        "priority": 50, "transform": "TRIM", "config": custom_config,
    }]
    with patch("backend.app.services.excel_rule_defaults._ensure_repeated_field_contracts"), \
         patch("backend.app.services.excel_rule_defaults._ensure_quantitation_impurity_name_contract"):
        ensure_excel_field_rules(database)

    database.save_system_field_rule.assert_not_called()
