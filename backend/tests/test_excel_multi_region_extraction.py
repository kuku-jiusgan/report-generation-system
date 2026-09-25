from pathlib import Path
import tempfile

from openpyxl import Workbook
import pytest

from backend.app.services.excel_field_extractor import extract_excel_fields
from backend.app.services.excel_catalog_manifest import prepared_fields, read_manifest
from backend.app.services.excel_rule_engine import ExcelRuleError


def test_repeat_block_concatenates_configured_regions_in_order() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    for row, values in enumerate((("A", 1), ("A", 2), ("B", 3), ("B", 4)), 1):
        sheet.cell(row, 1, values[0])
        sheet.cell(row, 2, values[1])
    field = {
        "fieldCode": "result.value", "cardinality": "MANY",
        "legacyJsonPath": "$.result[*].value",
    }
    rule = {
        "id": 1, "fieldCode": field["fieldCode"], "sourceType": "EXCEL",
        "enabled": True,
        "config": {
            "mode": "REPEAT_BLOCK", "sheet": "结果", "repeatCount": 1,
            "rowStart": 1, "rowEnd": 1, "startColumn": 1,
            "valueMode": "CELL", "sourcePath": field["legacyJsonPath"],
            "regions": [
                {"label": "A", "rowStart": 1, "rowEnd": 2, "startColumn": 2},
                {"label": "B", "rowStart": 3, "rowEnd": 4, "startColumn": 2},
            ],
        },
    }
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "regions.xlsx"
        workbook.save(path)

        payload = extract_excel_fields(path, [field], [rule])

    assert [record["value"] for record in payload["result"]] == ["1", "2", "3", "4"]


def test_excel_display_is_default_and_field_precision_overrides_it(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "结果"
    sheet["A1"] = 93.86666666666666
    sheet["A1"].number_format = "0.000"
    sheet["B1"] = 96.8
    sheet["B1"].number_format = "0.000"
    path = tmp_path / "display.xlsx"
    workbook.save(path)
    fields = [{"fieldCode": f"result.{name}", "cardinality": "ONE",
               "legacyJsonPath": f"$.result.{name}"}
              for name in ("raw", "fixed", "cell", "horizontal", "pairRaw", "pair")]
    base = {"mode": "FIXED_CELL", "sheet": "结果", "row": 1, "column": 1}
    repeated = {"mode": "REPEAT_BLOCK", "sheet": "结果", "rowStart": 1, "rowEnd": 1}
    rules = [{"fieldCode": f"result.{name}", "sourceType": "EXCEL", "enabled": True,
              "config": config} for name, config in (
                  ("raw", base),
                  ("fixed", {**base, "displayDecimals": 0}),
                  ("cell", {**repeated, "startColumn": 1}),
                  ("horizontal", {**repeated, "startColumn": 1, "valueMode": "HORIZONTAL_CELL",
                                   "valueCount": 1, "displayDecimals": 1}),
                  ("pairRaw", {**repeated, "valueMode": "CELL_PAIR", "pairColumns": [1, 2]}),
                  ("pair", {**repeated, "valueMode": "CELL_PAIR", "pairColumns": [1, 2],
                            "displayDecimals": 0}),
              )]

    payload = extract_excel_fields(path, fields, rules)

    assert payload["result"] == {"raw": "93.867", "fixed": "94", "cell": "93.867",
                                 "horizontal": "93.9", "pairRaw": "（93.867，96.800）",
                                 "pair": "（94，97）"}


def test_actual_workbook_extracts_injections_and_cached_statistics() -> None:
    root = Path(__file__).parents[2]
    manifest = read_manifest(root / "mapping/intermediate-precision-excel.json")
    group = manifest["groupCode"]
    fields, rules = [], []
    group_meta = {"groupCode": group, "cardinality": "MANY", "fields": [
        {"fieldCode": spec["code"], "label": spec.get("label") or spec["code"],
         "jsonKey": spec["code"].rsplit(".", 1)[-1]}
        for spec in manifest["fields"] if spec["code"].startswith("uncategorized.")
    ], "levels": [{"levelKey": "injections", "kind": "ARRAY"}]}
    for index, spec in enumerate(prepared_fields(manifest, group_meta), 1):
        code = spec["fieldCode"]
        path = spec["legacyJsonPath"]
        fields.append({"fieldCode": code, "groupCode": group, "cardinality": "MANY",
                       "legacyJsonPath": path})
        rules.append({"id": index, "fieldCode": code, "sourceType": "EXCEL", "enabled": True,
                      "config": spec["config"]})

    payload = extract_excel_fields(
        root / "文霞-标准计算表-动态-V56-解锁版-XM2026219-03.xlsm", fields, rules,
    )
    records = payload[group]
    assert len(records) == 3
    assert [item["impurityName"] for item in records] == [
        "3-吡啶磺酸甲酯", "3-吡啶磺酸乙酯", "3-吡啶磺酸异丙酯",
    ]
    assert [(records[0][key]["name"], records[0][key]["date"]) for key in ("technicianA", "technicianB")] == [
        ("技术员A", "2026.08.25"), ("技术员B", "2026.08.26")]
    assert all([len(record[key][detail]) for key, detail in (("technicianA", "injectionsA"), ("technicianB", "injectionsB"))] == [6, 6] for record in records)
    assert records[0]["technicianA"]["injectionsA"][0] == {
        "sequence": "1", "weighing": "13.74", "retentionTime": "4.208",
        "peakArea": "1177629", "concentration": "30.49", "relativeContent": "44.4",
    }
    assert records[0]["technicianB"]["injectionsB"][0] == {
        "sequence": "1", "weighing": "13.74", "retentionTime": "4.205",
        "peakArea": "766489", "concentration": "32.59", "relativeContent": "47.4",
    }
    assert [(records[0][key]["retentionRsd"], records[0][key]["contentRsd"]) for key in ("technicianA", "technicianB")] == [
        ("0.1", "2.8"), ("0.1", "1.1")]
    assert [records[0][key]["contentConfidenceInterval"] for key in ("technicianA", "technicianB")] == [
        "（42.700，45.300）", "（47.500，48.500）"]
    assert [records[0][key]["theoreticalPercentInterval"] for key in ("technicianA", "technicianB")] == [
        "（170.800，181.200）", "（190.000，194.000）"]
    assert records[0]["field_141"] == "5"
    assert [record["field_142"] for record in records] == [
        "（44.586，47.4645）", "（52.571，54.1125）", "（53.755，54.2113）",
    ]
    assert records[0]["field_143"] == "（178.342，189.858）"
    assert not payload["_meta"]["warnings"]


def test_technicians_are_declared_as_independent_groups() -> None:
    manifest_path = Path(__file__).parents[2] / "mapping/intermediate-precision-excel.json"
    manifest = read_manifest(manifest_path)
    levels = {item["levelKey"]: item for item in manifest["levels"]}
    for letter in ("A", "B"):
        assert levels[f"technician{letter}"]["kind"] == "OBJECT"
        assert levels[f"injections{letter}"]["parentLevelKey"] == f"technician{letter}"
    assert {item["level"] for item in manifest["fields"]} == {"", "technicianA", "injectionsA", "technicianB", "injectionsB"}


def test_required_join_source_fails_instead_of_silent_gap(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "人员"
    path = tmp_path / "missing.xlsx"
    workbook.save(path)
    field = {"fieldCode": "group.person", "cardinality": "MANY",
             "legacyJsonPath": "$.group[*].person"}
    rule = {"fieldCode": field["fieldCode"], "sourceType": "EXCEL", "enabled": True,
            "config": {"mode": "REPEAT_BLOCK", "sheet": "人员", "repeatCount": 1,
                       "rowStart": 1, "rowEnd": 1, "required": True,
                       "valueMode": "JOIN_CELLS", "valueSources": [{"sheet": "人员", "row": 1, "column": 1}]}}
    with pytest.raises(ExcelRuleError, match="字段 group.person 提取失败"):
        extract_excel_fields(path, [field], [rule])


def test_join_literal_label_does_not_need_an_excel_label(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "日期"
    sheet["A1"] = "2026.08.25"
    path = tmp_path / "date.xlsx"
    workbook.save(path)
    field = {"fieldCode": "group.person", "cardinality": "MANY",
             "legacyJsonPath": "$.group[*].person"}
    config = {"mode": "REPEAT_BLOCK", "sheet": "日期", "repeatCount": 1,
              "rowStart": 1, "rowEnd": 1, "required": True,
              "valueMode": "JOIN_CELLS", "joinSeparator": "\n",
              "valueSources": [{"literal": "技术员A"}, {"sheet": "日期", "row": 1, "column": 1}]}
    rule = {"fieldCode": field["fieldCode"], "sourceType": "EXCEL", "enabled": True,
            "config": config}

    payload = extract_excel_fields(path, [field], [rule])
    assert payload["group"][0]["person"] == "技术员A\n2026.08.25"

    config["valueSources"][0]["sheet"] = "日期"
    with pytest.raises(ExcelRuleError, match="固定文字来源只能配置非空 literal"):
        extract_excel_fields(path, [field], [rule])
